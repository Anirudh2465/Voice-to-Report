import io
from pathlib import Path
from jinja2 import Environment, FileSystemLoader
from weasyprint import HTML as WeasyHTML
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

# Template directory (relative to project root)
TEMPLATE_DIR = Path(__file__).parent.parent.parent.parent / "templates" / "standard"
jinja_env = Environment(loader=FileSystemLoader(str(TEMPLATE_DIR)))


def render_html_report(modality: str, report_data: dict) -> str:
    """Render a Jinja2 HTML template for the given modality and report data."""
    template_name = f"{modality.lower()}.html.j2"
    try:
        template = jinja_env.get_template(template_name)
    except Exception:
        # Fallback to base template
        template = jinja_env.get_template("base.html.j2")
    return template.render(**report_data)


def generate_pdf(modality: str, report_data: dict) -> bytes:
    """Render HTML template and convert to PDF using WeasyPrint."""
    html_content = render_html_report(modality, report_data)
    css_path = Path(__file__).parent.parent.parent.parent / "templates" / "styles" / "letterhead.css"
    pdf_bytes = WeasyHTML(string=html_content, base_url=str(TEMPLATE_DIR)).write_pdf(
        stylesheets=[str(css_path)] if css_path.exists() else []
    )
    return pdf_bytes


def generate_docx(report_data: dict) -> bytes:
    """Generate a formatted DOCX report from report_data dict."""
    doc = Document()

    # ── Header / Letterhead ────────────────────────────────────────────────
    header = doc.add_heading(report_data.get("institution", "Radiology Department"), level=1)
    header.alignment = WD_ALIGN_PARAGRAPH.CENTER

    sub = doc.add_paragraph(report_data.get("radiologist_name", "") + " | " + report_data.get("designation", ""))
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER

    doc.add_paragraph()  # spacer

    # ── Patient Info ────────────────────────────────────────────────────────
    info_table = doc.add_table(rows=2, cols=3)
    info_table.style = "Table Grid"
    cells = info_table.rows[0].cells
    cells[0].text = f"Patient: {report_data.get('patient_name', 'N/A')}"
    cells[1].text = f"ID: {report_data.get('patient_id', 'N/A')}"
    cells[2].text = f"Date: {report_data.get('date', '')}"
    cells2 = info_table.rows[1].cells
    cells2[0].text = f"Modality: {report_data.get('modality', '')}"
    cells2[1].text = f"Referring Dr: {report_data.get('referring_doctor', 'N/A')}"
    cells2[2].text = ""

    doc.add_paragraph()

    # ── Report Sections ─────────────────────────────────────────────────────
    content = report_data.get("report_content", {})
    sections_order = ["clinical_history", "technique", "findings", "impression", "recommendations"]
    section_labels = {
        "clinical_history": "Clinical History",
        "technique": "Technique",
        "findings": "Findings",
        "impression": "Impression",
        "recommendations": "Recommendations",
    }

    for key in sections_order:
        if key in content and content[key]:
            heading = doc.add_heading(section_labels.get(key, key.title()), level=2)
            para = doc.add_paragraph(content[key])
            para.style.font.size = Pt(11)

    doc.add_paragraph()

    # ── Signature ────────────────────────────────────────────────────────────
    sig = doc.add_paragraph()
    sig.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = sig.add_run(f"Dr. {report_data.get('radiologist_name', '')}")
    run.bold = True
    sig2 = doc.add_paragraph()
    sig2.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    sig2.add_run(report_data.get("designation", ""))

    # ── Save to bytes ────────────────────────────────────────────────────────
    buffer = io.BytesIO()
    doc.save(buffer)
    return buffer.getvalue()


def generate_txt(report_data: dict) -> bytes:
    """Generate a plain-text version of the report."""
    lines = [
        f"RADIOLOGY REPORT",
        f"{'=' * 60}",
        f"Facility: {report_data.get('institution', '')}",
        f"Radiologist: {report_data.get('radiologist_name', '')}",
        f"{'=' * 60}",
        f"Patient: {report_data.get('patient_name', 'N/A')}",
        f"Patient ID: {report_data.get('patient_id', 'N/A')}",
        f"Modality: {report_data.get('modality', '')}",
        f"Date: {report_data.get('date', '')}",
        f"{'=' * 60}",
    ]
    content = report_data.get("report_content", {})
    for key, label in [
        ("clinical_history", "CLINICAL HISTORY"),
        ("technique", "TECHNIQUE"),
        ("findings", "FINDINGS"),
        ("impression", "IMPRESSION"),
        ("recommendations", "RECOMMENDATIONS"),
    ]:
        if key in content and content[key]:
            lines.append(f"\n{label}:")
            lines.append(content[key])

    lines.append(f"\n{'=' * 60}")
    lines.append(f"Signed by: Dr. {report_data.get('radiologist_name', '')}")
    return "\n".join(lines).encode("utf-8")
